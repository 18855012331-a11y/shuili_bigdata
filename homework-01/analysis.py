"""垂线位移回归分析。运行：python analysis.py；输出位于 outputs/。"""
from pathlib import Path
import argparse
import json
import os
import tempfile
os.environ.setdefault("MPLCONFIGDIR", str(Path(tempfile.gettempdir()) / "dam-regression-mpl"))
os.environ.setdefault("XDG_CACHE_HOME", str(Path(tempfile.gettempdir()) / "dam-regression-cache"))
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from sklearn.pipeline import make_pipeline
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import LinearRegression
from sklearn.model_selection import TimeSeriesSplit
from sklearn.metrics import r2_score, mean_squared_error, mean_absolute_error
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

MODELS = {'M1': ['h'], 'M2': ['h', 'h2'],
          'M3': ['h', 'h2', 'sin1', 'cos1', 'sin2', 'cos2', 't']}

def features(d, origin, h0):
    t = (d['时间'] - origin).dt.total_seconds().to_numpy() / (86400 * 365.25)
    h = d['上游水位(m)'].to_numpy() - h0
    return pd.DataFrame(dict(h=h, h2=h*h, sin1=np.sin(2*np.pi*t),
        cos1=np.cos(2*np.pi*t), sin2=np.sin(4*np.pi*t), cos2=np.cos(4*np.pi*t), t=t))

def new_model():
    return make_pipeline(StandardScaler(), LinearRegression())

def scores(y, pred):
    return {'R2': float(r2_score(y, pred)),
            'RMSE_mm': float(np.sqrt(mean_squared_error(y, pred))),
            'MAE_mm': float(mean_absolute_error(y, pred))}

def main():
    root = Path(__file__).resolve().parent
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--data', type=Path, default=root/'data/垂线数据.xlsx')
    parser.add_argument('--out', type=Path, default=root/'outputs')
    args = parser.parse_args()
    out = args.out; out.mkdir(parents=True, exist_ok=True)
    book = pd.ExcelFile(args.data)
    metrics, cleaning, summaries = [], [], []
    for sheet in book.sheet_names:
        raw = pd.read_excel(book, sheet_name=sheet)
        d = raw[['时间', '顺河向位移(mm)', '上游水位(m)']].copy()
        d['时间'] = pd.to_datetime(d['时间'], errors='coerce')
        for col in d.columns[1:]:
            d[col] = pd.to_numeric(d[col], errors='coerce')
        d = d.replace([np.inf, -np.inf], np.nan).dropna()
        valid = len(d)
        # 同一时刻的重复观测取均值；不插值、不按结果删除异常值。
        d = d.groupby('时间', as_index=False).mean().sort_values('时间').reset_index(drop=True)
        if len(d) < 50:
            raise ValueError(f'{sheet}: 有效数据不足50条')
        n = int(len(d)*0.8)
        train, test = d.iloc[:n], d.iloc[n:]
        origin = train['时间'].iloc[0]
        h0 = float(train['上游水位(m)'].mean())
        X = features(d, origin, h0)
        y = d['顺河向位移(mm)'].to_numpy()
        cleaning.append(dict(station=sheet, raw=len(raw), invalid=len(raw)-valid,
            duplicate_extra=valid-len(d), clean=len(d), train=n, test=len(d)-n,
            start=str(d['时间'].min()), end=str(d['时间'].max()),
            test_start=str(test['时间'].iloc[0])))
        fitted = {}; cv_results = {}
        for name, cols in MODELS.items():
            fold_errors = []
            for a, b in TimeSeriesSplit(n_splits=4).split(X.iloc[:n]):
                model = new_model().fit(X.iloc[a][cols], y[a])
                fold_errors.append(float(np.sqrt(mean_squared_error(y[b], model.predict(X.iloc[b][cols])))))
            cv = float(np.mean(fold_errors)); cv_results[name] = cv
            model = new_model().fit(X.iloc[:n][cols], y[:n])
            pred = model.predict(X[cols]); fitted[name] = (model, pred)
            tr = scores(y[:n], pred[:n]); te = scores(y[n:], pred[n:])
            p = len(cols)
            adj = 1-(1-tr['R2'])*(n-1)/(n-p-1)
            metrics.append(dict(station=sheet, model=name, cv_RMSE_mm=cv,
                train_R2=tr['R2'], train_adjusted_R2=adj,
                test_R2=te['R2'], test_RMSE_mm=te['RMSE_mm'], test_MAE_mm=te['MAE_mm']))
        # 仅按训练期内的交叉验证选模型，测试集不参与选择。
        selected = min(cv_results, key=cv_results.get)
        model, pred = fitted[selected]; cols = MODELS[selected]
        scaler = model.named_steps['standardscaler']
        lr = model.named_steps['linearregression']
        coef = lr.coef_ / scaler.scale_
        intercept = float(lr.intercept_ - np.dot(coef, scaler.mean_))
        pd.DataFrame({'factor':['intercept']+cols,'coefficient':[intercept]+list(coef)}).to_csv(out/f'{sheet}_coefficients.csv',index=False,encoding='utf-8-sig')
        result = d.copy(); result['split'] = ['train']*n+['test']*(len(d)-n)
        for name, (_, yp) in fitted.items(): result[f'{name}_prediction_mm'] = yp
        result['selected_residual_mm'] = y-pred
        result.to_csv(out/f'{sheet}_predictions.csv',index=False,encoding='utf-8-sig')
        res = y[n:]-pred[n:]
        acf1 = float(np.corrcoef(res[:-1],res[1:])[0,1]) if np.std(res)>0 else None
        summaries.append(dict(station=sheet,selected=selected,h0=h0,origin=str(origin),
            intercept=intercept,coefficients=dict(zip(cols,map(float,coef))),
            residual_lag1=acf1,**scores(y[n:],pred[n:])))
        fig, axes = plt.subplots(3,1,figsize=(10,9),constrained_layout=True)
        axes[0].plot(d['时间'],y,lw=1,label='Observed')
        axes[0].plot(d['时间'],pred,lw=1,label=f'{selected} predicted')
        axes[0].axvline(test['时间'].iloc[0],color='black',ls='--',label='Test begins')
        axes[0].set(ylabel='Displacement (mm)',title=sheet);axes[0].legend()
        axes[1].scatter(y[n:],pred[n:],s=10)
        lo=min(y[n:].min(),pred[n:].min());hi=max(y[n:].max(),pred[n:].max())
        axes[1].plot([lo,hi],[lo,hi],'k--');axes[1].set(xlabel='Observed (mm)',ylabel='Test prediction (mm)')
        axes[2].plot(test['时间'],res,lw=1);axes[2].axhline(0,color='black',ls='--')
        axes[2].set(ylabel='Test residual (mm)',xlabel='Date')
        fig.savefig(out/f'{sheet}.png',dpi=160);plt.close(fig)
    m = pd.DataFrame(metrics); m.to_csv(out/'model_metrics.csv',index=False,encoding='utf-8-sig')
    pd.DataFrame(cleaning).to_csv(out/'data_quality.csv',index=False,encoding='utf-8-sig')
    (out/'model_summary.json').write_text(json.dumps(summaries,ensure_ascii=False,indent=2),encoding='utf-8')
    report(out,m,cleaning,summaries)
    print(m.to_string(index=False));print('\nSelected models:',[(s['station'],s['selected']) for s in summaries])
    print('Report:',out/'垂线数据回归分析报告.docx')

def report(out,m,cleaning,summaries):
    doc = Document()
    normal=doc.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    doc.add_heading('大坝安全监测垂线数据回归分析',0)
    doc.add_paragraph('基础档｜Python 可复现分析｜5个测点')
    doc.add_heading('1. 分析目标与数据处理',level=1)
    doc.add_paragraph('以顺河向位移（mm）为因变量，以上游水位和由日期构造的因子为自变量，比较不同复杂度的回归模型。分析覆盖PL13-1至PL13-5；所有指标和图表均由随附代码实际运行生成。未使用渗压数据，也未将附件示例程序当作必须执行的指令。')
    doc.add_paragraph('处理流程：解析日期与数值，剔除无效或非有限值，将同一时刻的重复记录按均值合并，再按时间升序排列。不填补缺测日期，不额外删除极端观测。重复记录取均值是本次处理约定，实际工程应核查其来源。')
    table=doc.add_table(rows=1, cols=6);table.style='Table Grid'
    for c,v in zip(table.rows[0].cells,['测点','原始条数','无效条数','重复多余条数','训练/测试','日期范围']):c.text=v
    for q in cleaning:
        vals=[q['station'],q['raw'],q['invalid'],q['duplicate_extra'],f"{q['train']}/{q['test']}",q['start'][:10]+'至'+q['end'][:10]]
        for c,v in zip(table.add_row().cells,vals):c.text=str(v)
    doc.add_heading('2. 自变量选择与模型理解',level=1)
    doc.add_paragraph('令 y 为位移，h=H−H₀，H₀为该测点完整训练集的平均水位；t为距训练集起始时刻的年数（每年365.25天）。所有模型包含截距和误差项ε。')
    for text in ['M1：y=β₀+β₁h+ε。表示水位与位移的线性关联，是基准模型。',
        'M2：y=β₀+β₁h+β₂h²+ε。二次项允许水位响应呈弯曲变化；水位中心化有助于改善数值条件，但不保证消除共线性。',
        'M3：y=β₀+β₁h+β₂h²+a₁sin(2πt)+b₁cos(2πt)+a₂sin(4πt)+b₂cos(4πt)+ct+ε。年周期和半年周期项用于近似季节影响，线性时间项表示长期趋势。没有实测温度，因此周期项不能等同于温度实测效应。']:
        doc.add_paragraph(text)
    doc.add_paragraph('这是对构造后的自变量做多元线性回归：虽然包含h²、正弦和余弦，模型对待估系数仍然线性。采用标准化加普通最小二乘；标准化仅在每次拟合的数据上学习。系数文件已换算回上述因子的原始尺度，可直接代入方程。时间和水位因子可能相关，单个系数不应被当作独立因果效应。')
    doc.add_heading('3. 验证方法与评价指标',level=1)
    doc.add_paragraph('每个测点前80%的记录用于训练，后20%作为独立测试期。训练期内部采用4折扩展窗口时间序列交叉验证，按平均验证RMSE最小选择模型；测试集只用于最终评价。每折标准化参数仅由该折训练样本确定。h的参考零点由完整训练期确定，只是基函数平移，不含测试期信息；含截距和完整一次、二次项时不改变拟合空间。')
    doc.add_paragraph('R²=1−Σ(y−ŷ)²/Σ(y−ȳ)²，表示相对于对应评价集均值基准的误差改善；测试R²可以为负，且不代表因果关系。调整R²=1−(1−R²)(n−1)/(n−p−1)，仅对训练集报告，p为非截距因子数。RMSE=√mean((y−ŷ)²)，MAE=mean(|y−ŷ|)，两者单位均为mm，越小越好。不同测点的位移变幅不同，不能只凭R²比较工程风险。')
    table=doc.add_table(rows=1,cols=7);table.style='Table Grid'
    headers=['测点','模型','CV RMSE','训练R²','训练调整R²','测试R²','测试RMSE']
    for c,v in zip(table.rows[0].cells,headers):c.text=v
    for _,r in m.iterrows():
        vals=[r.station,r.model]+[f'{r[k]:.4f}' for k in ['cv_RMSE_mm','train_R2','train_adjusted_R2','test_R2','test_RMSE_mm']]
        for c,v in zip(table.add_row().cells,vals):c.text=v
    doc.add_paragraph('完整MAE及未舍入指标见model_metrics.csv。表中同时展示各候选模型的测试指标供比较，但最终模型的选择严格依据训练期交叉验证。')
    doc.add_heading('4. 各测点结果与图形诊断',level=1)
    for s,q in zip(summaries,cleaning):
        doc.add_heading(s['station'],level=2)
        doc.add_paragraph(f"训练期内选择{s['selected']}。测试期从{q['test_start'][:10]}开始，R²={s['R2']:.4f}，RMSE={s['RMSE_mm']:.4f} mm，MAE={s['MAE_mm']:.4f} mm。H₀={s['h0']:.6f} m，t的起点为{s['origin']}。")
        equation=f"ŷ = {s['intercept']:.8g}"+''.join(f" {v:+.8g}×{k}" for k,v in s['coefficients'].items())
        doc.add_paragraph(equation)
        doc.add_picture(str(out/f"{s['station']}.png"),width=Inches(6.1))
        doc.add_paragraph('图从上到下为全时段实测与预测、测试期实测与预测散点、测试期残差。虚线分隔训练和测试。')
        doc.add_paragraph(f"测试残差相邻观测相关系数为{s['residual_lag1']:.3f}。由于采样并非严格等间隔，这不是固定1日滞后相关；明显的连续偏差说明模型仍可能遗漏动态因素，不能据此假定误差相互独立。")
    doc.add_heading('5. 讨论、理解与局限',level=1)
    best_counts=pd.Series([s['selected'] for s in summaries]).value_counts().to_dict()
    doc.add_paragraph('按交叉验证选出的模型分布为：'+str(best_counts)+'。模型增加因子后，训练拟合通常会改善，但未来时段的误差不一定减小，最终应关注时间外推能力。若测试R²较低或为负，应如实说明模型不足，不通过随机打乱时间或反复选择测试集来美化结果。')
    doc.add_paragraph('本分析是教学用途的统计建模，并非大坝安全鉴定。水位、季节项和长期趋势提供可解释的经验分解，但不能证明真实作用机制；未纳入实测温度、运行工况、测量基准变动和滞后效应。没有对高阶项进行无约束扩展，也未给出需要独立同分布误差假设的显著性结论。异常残差应结合现场记录复核。')
    doc.add_paragraph('学习体会（供结合个人经历修改）：回归分析不只是得到一个较高的决定系数，更重要的是解释因子来源、避免未来信息泄漏，并检查预测误差随时间的变化。复杂模型的价值应由验证结果支持，而不是由公式长度决定。')
    doc.add_heading('6. 运行与复现',level=1)
    doc.add_paragraph('使用Python 3.10及以上版本，在项目目录执行：python -m pip install -r requirements.txt，然后执行python analysis.py。脚本从data/垂线数据.xlsx读取数据，在outputs生成本报告、15组模型指标、5个测点的预测、系数及图表。也可使用--data和--out指定输入文件和输出目录。无需网络服务或密钥。')
    doc.save(out/'垂线数据回归分析报告.docx')

if __name__=='__main__':
    main()
